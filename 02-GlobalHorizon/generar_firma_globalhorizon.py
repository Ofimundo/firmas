import pandas as pd
import os
import sys
from openpyxl import load_workbook

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# Obtener ruta base del script y raíz del proyecto
base_path = os.path.dirname(os.path.abspath(sys.executable if getattr(sys, 'frozen', False) else __file__))
proyecto_root = os.path.dirname(base_path)

# === Rutas ===
#ruta_template = os.path.join(base_path, "firma1.html")
ruta_template = os.path.join(base_path, "diseño_globalhorizon_2.html")
ruta_datos = os.path.join(proyecto_root, "empleados.xlsx")
ruta_salida = os.path.join(proyecto_root, "05-firmas-generadas", "02-GlobalHorizonLatam")

os.makedirs(ruta_salida, exist_ok=True)

# Leer plantilla HTML
with open(ruta_template, "r", encoding="utf-8") as f:
    template = f.read()

# Leer datos con pandas (solo para lógica)
empleados = pd.read_excel(ruta_datos, sheet_name="GLOBAL")

if "procesado" not in empleados.columns:
    raise ValueError("La columna 'procesado' no existe en empleados.xlsx")

empleados["procesado"] = empleados["procesado"].astype(str).str.strip()

# Abrir workbook original (preserva formato)
wb = load_workbook(ruta_datos)
ws = wb["GLOBAL"] # si necesitas una hoja específica

# Obtener número de columna "procesado"
col_procesado_excel = empleados.columns.get_loc("procesado") + 1  # +1 porque Excel empieza en 1

for index, row in empleados.iterrows():

    if row["procesado"].lower() != "no":
        continue

    html_personal = template

    for col in empleados.columns:
        html_personal = html_personal.replace(f"{{{{{col}}}}}", str(row[col]))

    nombre_archivo = str(row["nombre_empleado"]).replace(" ", "_").lower()
    salida = os.path.join(ruta_salida, f"{nombre_archivo}.html")

    with open(salida, "w", encoding="utf-8") as f:
        f.write(html_personal)

    print(f"✅ Firma generada: {salida}")

    # === Actualizar SOLO la celda correspondiente ===
    fila_excel = index + 2  # +2 porque pandas empieza en 0 y Excel tiene encabezado en fila 1
    ws.cell(row=fila_excel, column=col_procesado_excel, value="Si")

    wb.save(ruta_datos)  # guarda sin destruir formato

print("\nProceso finalizado.")