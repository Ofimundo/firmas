import pandas as pd
import os
import sys
import base64
import mimetypes
from openpyxl import load_workbook

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# Obtener ruta base del script (01-Ofimundo/emergencia)
base_path = os.path.dirname(os.path.abspath(sys.executable if getattr(sys, 'frozen', False) else __file__))
# Ruta a 01-Ofimundo
ofimundo_dir = os.path.dirname(base_path)
# Ruta a la raíz del proyecto (PROYECTO FIRMAS)
proyecto_root = os.path.dirname(ofimundo_dir)

# === Rutas del proyecto ===
ruta_template = os.path.join(ofimundo_dir, "diseño_ofimundo.html")
ruta_datos = os.path.join(base_path, "empleados.xlsx")

# Si el Excel de emergencia no existe, crear una copia inicial desde el original sin mezclar
if not os.path.exists(ruta_datos):
    ruta_datos_original = os.path.join(proyecto_root, "empleados.xlsx")
    if os.path.exists(ruta_datos_original):
        import shutil
        shutil.copy2(ruta_datos_original, ruta_datos)
        print(f"📋 Se creó una copia independiente de 'empleados.xlsx' en: {ruta_datos}")

ruta_salida = base_path

# Directorios de emergencia para imágenes
dir_foto_perfil = os.path.join(base_path, "foto perfil")
dir_qr_linkedin = os.path.join(base_path, "qr linkedin")
dir_qr_social = os.path.join(base_path, "qr social")

os.makedirs(ruta_salida, exist_ok=True)

def buscar_imagen_emergencia(directorio, nombre_corto, nombre_archivo, url_original):
    """
    Busca si existe la imagen en la carpeta de emergencia.
    Prioridades de búsqueda:
    1. <nombre_corto>.<ext>
    2. <nombre_archivo>.<ext>
    3. <basename_url>
    """
    if not os.path.exists(directorio):
        return None

    extensiones = [".png", ".jpg", ".jpeg", ".webp", ".gif"]
    nombres_base = []
    
    if nombre_corto and str(nombre_corto) != "nan":
        nombres_base.append(str(nombre_corto).strip())
    if nombre_archivo:
        nombres_base.append(str(nombre_archivo).strip())
    if url_original and str(url_original) != "nan":
        filename_url = os.path.basename(str(url_original).split("?")[0])
        name_without_ext = os.path.splitext(filename_url)[0]
        if name_without_ext:
            nombres_base.append(name_without_ext)
            nombres_base.append(filename_url)

    # 1. Probar combinaciones exactas de nombre y extensión
    for nb in nombres_base:
        if "." in nb:
            posible_path = os.path.join(directorio, nb)
            if os.path.isfile(posible_path):
                return os.path.abspath(posible_path).replace("\\", "/")
        for ext in extensiones:
            posible_path = os.path.join(directorio, f"{nb}{ext}")
            if os.path.isfile(posible_path):
                return os.path.abspath(posible_path).replace("\\", "/")
    
    # 2. Búsqueda insensible a mayúsculas/minúsculas en el directorio
    archivos_dir = os.listdir(directorio)
    for nb in nombres_base:
        nb_clean = os.path.splitext(nb)[0].lower()
        for f in archivos_dir:
            f_name_without_ext = os.path.splitext(f)[0].lower()
            if f_name_without_ext == nb_clean:
                return os.path.abspath(os.path.join(directorio, f)).replace("\\", "/")

def imagen_a_base64(ruta_imagen):
    """Convierte una imagen local a Data URI Base64 para incrustarla directamente en la firma HTML."""
    if not ruta_imagen or not os.path.isfile(ruta_imagen):
        return ruta_imagen
    mime_type, _ = mimetypes.guess_type(ruta_imagen)
    if not mime_type:
        mime_type = "image/png"
    with open(ruta_imagen, "rb") as img_file:
        encoded = base64.b64encode(img_file.read()).decode("utf-8")
    return f"data:{mime_type};base64,{encoded}"

# Leer plantilla HTML
with open(ruta_template, "r", encoding="utf-8") as f:
    template = f.read()

# Leer datos
empleados = pd.read_excel(ruta_datos, sheet_name="OFIMUNDO")

if "procesado" not in empleados.columns:
    raise ValueError("La columna 'procesado' no existe en empleados.xlsx")

empleados["procesado"] = empleados["procesado"].astype(str).str.strip()

# Abrir workbook para actualización
wb = load_workbook(ruta_datos)
ws = wb["OFIMUNDO"]

col_procesado_excel = empleados.columns.get_loc("procesado") + 1

firmas_generadas = 0

for index, row in empleados.iterrows():

    if row["procesado"].lower() != "no":
        continue

    firmas_generadas += 1
    html_personal = template

    nombre_empleado_clean = str(row["nombre_empleado"]).strip()
    nombre_archivo = nombre_empleado_clean.replace(" ", "_").lower()
    nombre_corto = str(row.get("nombre_corto", "")).strip()

    row_data = row.to_dict()

    # Sustitución de imágenes por versión de emergencia si existe en las carpetas locales
    foto_emergencia = buscar_imagen_emergencia(dir_foto_perfil, nombre_corto, nombre_archivo, row.get("foto_empleado"))
    if foto_emergencia:
        row_data["foto_empleado"] = imagen_a_base64(foto_emergencia)
        print(f"  📷 [EMERGENCIA] Foto de perfil local incrustada (Base64): {foto_emergencia}")

    qr_linkedin_emergencia = buscar_imagen_emergencia(dir_qr_linkedin, nombre_corto, nombre_archivo, row.get("qr_linkedin"))
    if qr_linkedin_emergencia:
        row_data["qr_linkedin"] = imagen_a_base64(qr_linkedin_emergencia)
        print(f"  🔗 [EMERGENCIA] QR LinkedIn local incrustado (Base64): {qr_linkedin_emergencia}")

    qr_social_emergencia = buscar_imagen_emergencia(dir_qr_social, nombre_corto, nombre_archivo, row.get("qr_social"))
    if qr_social_emergencia:
        row_data["qr_social"] = imagen_a_base64(qr_social_emergencia)
        print(f"  📲 [EMERGENCIA] QR Social local incrustado (Base64): {qr_social_emergencia}")

    for col, val in row_data.items():
        val_str = str(val).strip() if pd.notna(val) else ""
        html_personal = html_personal.replace(f"{{{{{col}}}}}", val_str)

    salida = os.path.join(ruta_salida, f"{nombre_archivo}.html")

    with open(salida, "w", encoding="utf-8") as f:
        f.write(html_personal)

    print(f"✅ Firma generada (modo emergencia): {salida}")

    fila_excel = index + 2
    ws.cell(row=fila_excel, column=col_procesado_excel, value="Si")

    try:
        wb.save(ruta_datos)
    except PermissionError:
        print("⚠️ ALERTA: No se pudo actualizar 'empleados.xlsx' porque el archivo está abierto en Microsoft Excel. Por favor, ciérralo y vuelve a ejecutar el script.")
        break

if firmas_generadas == 0:
    print("ℹ️ No se generó ninguna firma porque todas las filas en el Excel tienen 'procesado' = 'Si'.")
    print("   (Para probar, cambia 'procesado' a 'No' en empleados.xlsx para la persona correspondiente).")

print("\nProceso finalizado.")
