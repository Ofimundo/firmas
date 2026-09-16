import pandas as pd
import os
import sys
from openpyxl import load_workbook

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# Obtener ruta base del script o ejecutable y raíz del proyecto
base_path = os.path.dirname(os.path.abspath(sys.executable if getattr(sys, 'frozen', False) else __file__))
proyecto_root = os.path.dirname(base_path)

# === Rutas ===
ruta_template = os.path.join(base_path, "diseño_ofimundo.html")
ruta_datos = os.path.join(proyecto_root, "empleados.xlsx")
ruta_salida = os.path.join(proyecto_root, "05-firmas-generadas", "01-Ofimundo")

os.makedirs(ruta_salida, exist_ok=True)

# Leer plantilla HTML
with open(ruta_template, "r", encoding="utf-8") as f:
    template = f.read()

# Leer datos
empleados = pd.read_excel(ruta_datos, sheet_name="OFIMUNDO")

if "procesado" not in empleados.columns:
    raise ValueError("La columna 'procesado' no existe en empleados.xlsx")

empleados["procesado"] = empleados["procesado"].astype(str).str.strip()

# Abrir workbook
wb = load_workbook(ruta_datos)
ws = wb["OFIMUNDO"] # si necesitas una hoja específica

# Obtener número de columna "procesado"
col_procesado_excel = empleados.columns.get_loc("procesado") + 1

for index, row in empleados.iterrows():

    if row["procesado"].lower() != "no":
        continue

    html_personal = template

    for col in empleados.columns:
        val = str(row[col]).strip() if pd.notna(row[col]) else ""
        html_personal = html_personal.replace(f"{{{{{col}}}}}", val)

    nombre_empleado_clean = str(row["nombre_empleado"]).strip()
    nombre_archivo = nombre_empleado_clean.replace(" ", "_").lower()
    salida = os.path.join(ruta_salida, f"{nombre_archivo}.html")

    with open(salida, "w", encoding="utf-8") as f:
        f.write(html_personal)

    print(f"✅ Firma generada: {salida}")

    # === Actualizar SOLO la celda correspondiente ===
    fila_excel = index + 2
    ws.cell(row=fila_excel, column=col_procesado_excel, value="Si")

    try:
        wb.save(ruta_datos)
    except PermissionError:
        print("⚠️ ALERTA: No se pudo actualizar 'empleados.xlsx' porque el archivo está abierto en Microsoft Excel. Por favor, ciérralo y vuelve a ejecutar el script.")
        break

print("\nProceso finalizado.")